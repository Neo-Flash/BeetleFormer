"""Check original external sequence provenance against published source tables.

Source evidence does not change the frozen sequences, labels or model inputs.
Gaddelapati genomic fragments retain the original T7/leading-G trimming rule;
the full source primer strings are recorded to make that derivation explicit.
"""
import hashlib
import json
import re
from pathlib import Path

import pandas as pd
from docx import Document

from common import ROOT


def rc(seq):
    return seq.translate(str.maketrans('ACGT', 'TGCA'))[::-1]


def main():
    folder = ROOT / 'data/raw/original_external/source_evidence'
    ext = pd.read_csv(ROOT / 'provenance/original_external_records.csv')
    document = Document(folder / '42003_2024_6212_MOESM1_ESM.docx')
    symbols, primers = {}, {}
    for table in document.tables[:3]:
        for row in table.rows[1:]:
            match = re.search(r'TC\d+', row.cells[2].text)
            if match:
                symbols[match[0]] = row.cells[0].text.strip()
    for row in document.tables[4].rows:
        name = re.sub(r'\s+', '', row.cells[0].text)
        match = re.fullmatch(r'ds(TC\d+)([FR])', name)
        if match:
            primer = re.sub(r'\s+', '', row.cells[1].text).upper()
            assert primer.startswith('TAATACGACTCACTATAG')
            primers[(match[1], match[2])] = primer
    cedden = pd.read_excel(folder / '12915_2025_2219_MOESM1_ESM.xlsx',
                           sheet_name='Primer list 2', header=None)
    knorr = pd.read_excel(folder / '41598_2018_20416_MOESM1_ESM.xlsx',
                          sheet_name='S Table 1 - Potential targets', header=None)
    records = []
    for row in ext[ext.label.eq(1)].itertuples():
        record = dict(id=row.id, source=row.source_paper, source_gene=row.gene,
                      sequence_sha256=hashlib.sha256(row.sequence.encode()).hexdigest())
        if row.source_paper == 'Gaddelapati2024':
            record['published_gene_symbol'] = symbols[row.gene]
            fwd, rev = (primers[(row.gene, direction)] for direction in ['F', 'R'])
            trimmed = lambda p: p[len('TAATACGACTCACTATAG'):].lstrip('G')
            assert row.sequence.startswith(trimmed(fwd))
            assert row.sequence.endswith(rc(trimmed(rev)))
            record.update(source_table='Supplementary Tables 1–3 and 5',
                          published_forward_primer=fwd, published_reverse_primer=rev,
                          check='Original genomic fragment boundaries follow the archived T7/leading-G trimming rule; target identity matches the source gene table.')
        elif row.source_paper == 'Cedden2025':
            matches = cedden[cedden[5].astype(str).str.strip().eq(row.sequence)]
            assert len(matches) >= 1
            assert all(row.gene in str(v) for v in matches[0])
            record.update(source_table='Additional file 1, Primer list 2',
                          check='Full original sequence exactly matches the published amplicon sequence.')
        elif row.source_paper == 'Knorr2018':
            matches = knorr[knorr.apply(lambda c: c.astype(str).str.strip().eq(row.sequence)).any(axis=1)]
            assert len(matches) == 1
            assert row.gene in matches.iloc[0].astype(str).tolist()
            record.update(source_table='Supplementary Table 1',
                          check='Full original sequence and Arp1 gene identifier match the published table.')
        elif row.source_paper == 'AbdElHalim2016':
            fwd, rev = 'AAGGCAAGGACATATTCCGA', 'TCAAACGTGTAGATGCCAGTG'
            assert len(row.sequence) == 217
            assert row.sequence.startswith(fwd) and row.sequence.endswith(rc(rev))
            record.update(source_table='Table 1; https://pmc.ncbi.nlm.nih.gov/articles/PMC4944135/',
                          published_forward_primer=fwd, published_reverse_primer=rev,
                          check='Original 217-base fragment matches both published PCR primer boundaries.')
        else:
            raise ValueError(row.source_paper)
        records.append(record)
    assert len(records) == 20
    report = dict(positive_records_checked=20, sequences_or_labels_changed=False,
                  source_file_sha256={p.name: hashlib.sha256(p.read_bytes()).hexdigest()
                                      for p in folder.iterdir() if p.is_file()},
                  records=records)
    (ROOT / 'provenance/original_external_source_checks.json').write_text(
        json.dumps(report, ensure_ascii=False, indent=2))
    print('Published source association checks passed for all 20 original external positives.')


if __name__ == '__main__':
    main()
